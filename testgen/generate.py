"""
Report generation module - refactored from generate_report.py
"""
import json
import os
import glob
from datetime import datetime
from pathlib import Path
from typing import List, Tuple, Dict, Any

try:
    import openpyxl
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

from .config import Config
from .utils import ensure_dir_exists, format_error, format_success


def merge_batches(batches_dir: str) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Merges response.json files from batch directories.
    
    Returns:
        Tuple of (test_cases, errors)
    """
    all_test_cases = []
    errors = []
    
    # Find all batch folders
    search_path = os.path.join(batches_dir, "batch_*")
    batch_folders = sorted(glob.glob(search_path))
    
    if not batch_folders:
        print(f"❌ No 'batch_*' folders found in {batches_dir}")
        return [], []
    
    print(f"Found {len(batch_folders)} batch folders.")
    
    for folder in batch_folders:
        if not os.path.isdir(folder):
            continue
        
        json_path = os.path.join(folder, "response.json")
        folder_name = os.path.basename(folder)
        
        if os.path.exists(json_path):
            # Check if file is empty
            if os.path.getsize(json_path) == 0:
                error_msg = f"{folder_name}: response.json is EMPTY"
                print(f"  ❌ {error_msg}")
                errors.append(error_msg)
                continue
            
            try:
                with open(json_path, 'r') as f:
                    data = json.load(f)
                    if "test_cases" in data:
                        cases = data.get("test_cases", [])
                        print(f"  ✓ {folder_name}: Loaded {len(cases)} test cases.")
                        all_test_cases.extend(cases)
                    else:
                        error_msg = f"{folder_name}: 'test_cases' key not found in JSON"
                        print(f"  ❌ {error_msg}")
                        errors.append(error_msg)
            except json.JSONDecodeError as e:
                error_msg = f"{folder_name}: Invalid JSON - {e}"
                print(f"  ❌ {error_msg}")
                errors.append(error_msg)
            except Exception as e:
                error_msg = f"{folder_name}: Error reading file - {e}"
                print(f"  ❌ {error_msg}")
                errors.append(error_msg)
        else:
            error_msg = f"{folder_name}: response.json not found"
            print(f"  ❌ {error_msg}")
            errors.append(error_msg)
    
    return all_test_cases, errors


def generate_excel(test_cases: List[Dict], output_dir: str, template_path: str) -> bool:
    """
    Generates the Excel report from test cases.
    
    Returns:
        True if successful, False otherwise
    """
    if openpyxl is None:
        print(format_error("'openpyxl' library is required."))
        print("Install it using: pip install openpyxl")
        return False
    
    if not test_cases:
        print("No test cases to write.")
        return False
    
    # Generate filename with current date
    current_date = datetime.now().strftime("%Y-%m-%d")
    output_filename = f"Generated_Test_Cases_{current_date}.xlsx"
    output_file = os.path.join(output_dir, output_filename)
    
    if not os.path.exists(template_path):
        print(format_error(f"Template file not found: {template_path}"))
        return False
    
    # Ensure output directory exists
    ensure_dir_exists(output_dir)
    
    try:
        print(f"Generating report: {output_file}...")
        
        # Load Excel Template
        wb = openpyxl.load_workbook(template_path)
        
        if "Page1" in wb.sheetnames:
            ws = wb["Page1"]
        else:
            ws = wb.active
        
        # Clear existing data rows (keep header at row 1)
        max_row = ws.max_row
        if max_row > 1:
            for row in ws.iter_rows(min_row=2, max_row=max_row):
                for cell in row:
                    cell.value = None
                    cell.style = 'Normal'
        
        # Write new data starting at row 2
        start_row = 2
        
        # Define column widths
        col_widths = {
            1: 15,  # ID
            2: 10,  # Priority
            3: 30,  # Title
            4: 25,  # Precondition
            5: 50,  # Steps
            6: 40,  # Expected
            7: 30,  # Actual
            8: 10,  # Status
            9: 15,  # Tester
            10: 20  # Notes
        }
        
        # Apply column widths
        for col_idx, width in col_widths.items():
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width
        
        for i, case in enumerate(test_cases):
            row_num = start_row + i
            
            def get_val(key):
                val = case.get(key, "")
                if val is None:
                    return ""
                return str(val)
            
            # Map data to columns (1-based index)
            columns_data = [
                (1, "id"),
                (2, "priority"),
                (3, "title"),
                (4, "precondition"),
                (5, "steps"),
                (6, "expected_result"),
                (7, "actual_result"),
                (8, "status"),
                (9, "qa_tester"),
                (10, "qa_notes")
            ]
            
            max_lines_in_row = 1
            
            for col_idx, key in columns_data:
                cell = ws.cell(row=row_num, column=col_idx)
                value = get_val(key)
                cell.value = value
                
                # Apply text wrapping for long fields
                if col_idx in [3, 4, 5, 6, 7, 10]:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                    
                    # Estimate row height
                    newlines = value.count('\n')
                    char_limit = col_widths.get(col_idx, 20)
                    
                    if len(value) > char_limit:
                        wrapped_lines = (len(value) / char_limit)
                        lines = newlines + wrapped_lines + 1
                    else:
                        lines = newlines + 1
                    
                    if lines > max_lines_in_row:
                        max_lines_in_row = lines
                else:
                    cell.alignment = Alignment(vertical='top')
            
            # Adjust Row Height
            if max_lines_in_row > 1:
                height = min(max_lines_in_row * 15, 409)
                ws.row_dimensions[row_num].height = height
        
        # Save output
        wb.save(output_file)
        print(format_success(f"Successfully generated report at: {output_file}"))
        return True
        
    except Exception as e:
        print(format_error(f"Error during Excel generation: {e}"))
        import traceback
        traceback.print_exc()
        return False


def generate_report(config: Config, batches_dir: Optional[str] = None, output_dir: Optional[str] = None) -> bool:
    """
    Main function to generate test case report.
    
    Args:
        config: Configuration object
        batches_dir: Override batches directory
        output_dir: Override output directory
    
    Returns:
        True if successful, False otherwise
    """
    # Get configuration
    if batches_dir is None:
        batches_dir = config.resolve_path(config.get('directories.batches', './batches'))
    
    if output_dir is None:
        output_dir = config.resolve_path(config.get('directories.output', './output'))
    
    template_path = config.get_template_path()
    
    if not template_path:
        print(format_error("Excel template not found."))
        print("Please configure template path using:")
        print("  testgen config set template.primary /path/to/template.xlsx")
        return False
    
    print("--- Starting Test Case Report Generation ---")
    print(f"Searching for batches in: {batches_dir}")
    
    # Step 1: Merge
    test_cases, errors = merge_batches(batches_dir)
    
    # Stop if there are any errors
    if errors:
        print("\n" + "=" * 60)
        print("⚠️  ERROR: Cannot generate report due to batch issues:")
        print("=" * 60)
        for i, error in enumerate(errors, 1):
            print(f"{i}. {error}")
        print("\nPlease fix the above issues and run the script again.")
        print("=" * 60)
        return False
    
    if not test_cases:
        print("No test cases found. Exiting.")
        return False
    
    print(f"\nTotal merged test cases: {len(test_cases)}")
    
    # Step 2: Generate Excel
    success = generate_excel(test_cases, output_dir, template_path)
    
    if success:
        print("--- Done ---")
    
    return success
